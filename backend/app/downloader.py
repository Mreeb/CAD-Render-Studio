import os
import re
import time
import threading
from io import BytesIO
from typing import Optional, Callable, Dict, Any
import openpyxl
import requests
from requests.adapters import HTTPAdapter
from urllib3.util import Retry
from PIL import Image

from backend.app.config import settings
from backend.app.db import get_db

def extract_first_serial(raw_val: Any) -> str:
    if raw_val is None:
        return ""
    val_str = str(raw_val).strip().replace('\xa0', '')
    if not val_str:
        return ""
    parts = re.split(r'[,;/|\r\n]', val_str)
    for p in parts:
        cleaned = p.strip()
        if cleaned:
            return cleaned
    return ""

def is_cell_blank(val: Any) -> bool:
    if val is None:
        return True
    val_str = str(val).strip().replace('\xa0', '')
    return val_str == ""

def build_robust_session() -> requests.Session:
    session = requests.Session()
    retries = Retry(
        total=4,
        backoff_factor=1.0,
        status_forcelist=[429, 500, 502, 503, 504],
        raise_on_status=False
    )
    adapter = HTTPAdapter(max_retries=retries)
    session.mount('https://', adapter)
    session.mount('http://', adapter)
    return session

def get_product_image_url(serial_number: str, session: requests.Session, timeout: int = 25) -> Optional[str]:
    api_headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'application/json, text/plain, */*'
    }
    # 1. SuiteCommerce API
    try:
        api_url = f"https://www.auveco.com/api/items?q={serial_number}&fieldset=details"
        r = session.get(api_url, headers=api_headers, timeout=timeout)
        if r.status_code == 200:
            data = r.json()
            items = data.get("items", [])
            if items:
                first_item = items[0]
                img_detail = first_item.get("itemimages_detail", {})
                if isinstance(img_detail, dict):
                    for key, media_group in img_detail.items():
                        if isinstance(media_group, dict) and "urls" in media_group:
                            urls = media_group["urls"]
                            if urls and isinstance(urls[0], dict) and "url" in urls[0]:
                                candidate_url = urls[0]["url"]
                                if candidate_url and "no_image_available" not in candidate_url and "logo" not in candidate_url:
                                    return candidate_url
    except Exception:
        pass

    # 2. Probe Fallback
    probe_headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8'
    }
    probes = [
        f"https://www.auveco.com/site/item-images/{serial_number}_image-1.jpg",
        f"https://www.auveco.com/site/item-images/{serial_number}.jpg",
        f"https://www.auveco.com/site/item-images/{serial_number}_1.jpg"
    ]
    for p_url in probes:
        try:
            head_r = session.head(p_url, headers=probe_headers, timeout=8)
            if head_r.status_code == 200 and "image" in head_r.headers.get("Content-Type", ""):
                return p_url
        except Exception:
            pass

    return None

def download_and_save_png(image_url: str, dest_png_path: str, session: requests.Session, max_retries: int = 3, timeout: int = 25) -> bool:
    img_headers = {
        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
        'Accept': 'image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8'
    }
    tmp_path = dest_png_path + ".tmp"

    for attempt in range(1, max_retries + 1):
        try:
            r = session.get(image_url, headers=img_headers, timeout=timeout)
            if r.status_code == 200 and len(r.content) > 0:
                img = Image.open(BytesIO(r.content))
                if img.mode in ("P", "LA"):
                    img = img.convert("RGBA")
                elif img.mode not in ("RGB", "RGBA"):
                    img = img.convert("RGB")
                
                img.save(tmp_path, format="PNG")
                
                if os.path.exists(dest_png_path):
                    os.remove(dest_png_path)
                os.rename(tmp_path, dest_png_path)
                return True
        except Exception:
            if attempt < max_retries:
                time.sleep(1.0 * attempt)

    if os.path.exists(tmp_path):
        try: os.remove(tmp_path)
        except OSError: pass
    return False

class DownloaderService:
    _instance = None
    _lock = threading.Lock()

    def __new__(cls):
        with cls._lock:
            if cls._instance is None:
                cls._instance = super(DownloaderService, cls).__new__(cls)
                cls._instance._init_service()
            return cls._instance

    def _init_service(self):
        self.is_running = False
        self.stop_requested = False
        self.thread: Optional[threading.Thread] = None
        self.on_progress_callback: Optional[Callable[[Dict[str, Any]], None]] = None

    def get_state(self) -> Dict[str, Any]:
        with get_db() as conn:
            row = conn.execute("SELECT * FROM downloader_state WHERE id = 1").fetchone()
            if row:
                return dict(row)
            return {"status": "idle"}

    def update_state(self, updates: Dict[str, Any]):
        fields = ", ".join([f"{k} = ?" for k in updates.keys()])
        values = list(updates.values())
        with get_db() as conn:
            conn.execute(f"UPDATE downloader_state SET {fields}, updated_at = CURRENT_TIMESTAMP WHERE id = 1", values)
        
        # Broadcast via callback if listening
        if self.on_progress_callback:
            try:
                state = self.get_state()
                self.on_progress_callback(state)
            except Exception:
                pass

    def start_download(self, max_downloads: int = 100, delay: float = 0.5, timeout: int = 25) -> Dict[str, Any]:
        with self._lock:
            if self.is_running:
                raise RuntimeError("AuVeCo Image Downloader is already running.")

            self.is_running = True
            self.stop_requested = False

            self.update_state({
                "status": "running",
                "total_scanned": 0,
                "eligible_found": 0,
                "unique_serials": 0,
                "already_present": 0,
                "duplicates_skipped": 0,
                "downloaded_count": 0,
                "not_found_count": 0,
                "failed_count": 0,
                "current_serial": "Initializing...",
                "max_downloads": max_downloads,
                "started_at": time.strftime("%Y-%m-%d %H:%M:%S")
            })

            self.thread = threading.Thread(
                target=self._run_downloader,
                args=(max_downloads, delay, timeout),
                daemon=True
            )
            self.thread.start()
            return self.get_state()

    def stop_download(self) -> Dict[str, Any]:
        with self._lock:
            if not self.is_running:
                return self.get_state()
            self.stop_requested = True
            return self.get_state()

    def _run_downloader(self, max_downloads: int, delay: float, timeout: int):
        excel_path = settings.EXCEL_FILE
        output_dir = settings.DATA_DIRECTORY
        os.makedirs(output_dir, exist_ok=True)

        if not os.path.exists(excel_path):
            self.update_state({
                "status": "failed",
                "current_serial": f"Error: Excel file missing at {excel_path}"
            })
            self.is_running = False
            return

        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            ws = wb.active

            auveco_col = None
            phantom_col = None
            for col_idx in range(1, ws.max_column + 1):
                header_val = str(ws.cell(1, col_idx).value or '').strip().lower()
                if 'auveco' in header_val and auveco_col is None:
                    auveco_col = col_idx
                elif 'phantom' in header_val and phantom_col is None:
                    phantom_col = col_idx

            if auveco_col is None: auveco_col = 1
            if phantom_col is None: phantom_col = 2

            total_scanned = 0
            eligible_found = 0
            unique_serials = set()
            already_present = 0
            duplicates_skipped = 0
            downloaded_count = 0
            not_found_count = 0
            failed_count = 0

            session = build_robust_session()
            seen_serials_in_run = set()

            for r in range(2, ws.max_row + 1):
                if self.stop_requested:
                    self.update_state({"status": "stopped", "current_serial": "Stopped by user"})
                    break

                total_scanned += 1
                val_a = ws.cell(r, auveco_col).value
                val_b = ws.cell(r, phantom_col).value

                if is_cell_blank(val_a) or is_cell_blank(val_b):
                    continue

                eligible_found += 1
                raw_auveco_str = str(val_a).strip().replace('\xa0', '')
                first_sn = extract_first_serial(raw_auveco_str)

                if not first_sn:
                    continue

                unique_serials.add(first_sn)
                target_png_path = os.path.join(output_dir, f"{first_sn}.png")

                # Check if file exists on disk
                if os.path.exists(target_png_path) and os.path.getsize(target_png_path) > 0:
                    already_present += 1
                    seen_serials_in_run.add(first_sn)
                    self.update_state({
                        "total_scanned": total_scanned,
                        "eligible_found": eligible_found,
                        "unique_serials": len(unique_serials),
                        "already_present": already_present,
                        "current_serial": first_sn
                    })
                    continue

                # Check if duplicate in current run
                if first_sn in seen_serials_in_run:
                    duplicates_skipped += 1
                    self.update_state({
                        "total_scanned": total_scanned,
                        "eligible_found": eligible_found,
                        "duplicates_skipped": duplicates_skipped,
                        "current_serial": first_sn
                    })
                    continue

                seen_serials_in_run.add(first_sn)

                if downloaded_count >= max_downloads:
                    break

                self.update_state({
                    "total_scanned": total_scanned,
                    "eligible_found": eligible_found,
                    "unique_serials": len(unique_serials),
                    "already_present": already_present,
                    "duplicates_skipped": duplicates_skipped,
                    "downloaded_count": downloaded_count,
                    "not_found_count": not_found_count,
                    "failed_count": failed_count,
                    "current_serial": f"Searching: {first_sn}"
                })

                # Discover URL
                image_url = get_product_image_url(first_sn, session, timeout=timeout)
                if not image_url:
                    not_found_count += 1
                    self.update_state({"not_found_count": not_found_count, "current_serial": f"Not found: {first_sn}"})
                    continue

                # Download PNG
                self.update_state({"current_serial": f"Downloading: {first_sn}.png"})
                success = download_and_save_png(image_url, target_png_path, session, max_retries=3, timeout=timeout)
                if success:
                    downloaded_count += 1
                    self.update_state({"downloaded_count": downloaded_count, "current_serial": f"Saved: {first_sn}.png"})
                else:
                    failed_count += 1
                    self.update_state({"failed_count": failed_count, "current_serial": f"Failed: {first_sn}"})

                if delay > 0:
                    time.sleep(delay)

            if not self.stop_requested:
                self.update_state({
                    "status": "completed",
                    "current_serial": "Download process completed successfully"
                })

            # Auto-trigger reconciliation after download finishes
            from backend.app.reconciler import reconcile_database_and_filesystem
            reconcile_database_and_filesystem()

        except Exception as e:
            self.update_state({
                "status": "failed",
                "current_serial": f"Error: {str(e)}"
            })
        finally:
            self.is_running = False

downloader_service = DownloaderService()
